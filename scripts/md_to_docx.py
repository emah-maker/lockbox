#!/usr/bin/env python
"""Minimal Markdown -> DOCX converter (python-docx).

Handles: ATX headings, GitHub tables, bullet/numbered lists, fenced code
blocks, horizontal rules, and inline **bold**, `code`, and [text](url)
hyperlinks. Good enough for FRAIM narrative deliverables when pandoc is
unavailable.

Usage: python scripts/md_to_docx.py <input.md> <output.docx>
"""
import re
import sys

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.shared import Pt, RGBColor

LINK = re.compile(r"\[([^\]]+)\]\(([^)]+)\)")
INLINE = re.compile(r"(\*\*.+?\*\*|`[^`]+`|\[[^\]]+\]\([^)]+\))")


def add_hyperlink(paragraph, url, text):
    part = paragraph.part
    r_id = part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), r_id)
    new_run = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "0563C1")
    rpr.append(color)
    u = OxmlElement("w:u")
    u.set(qn("w:val"), "single")
    rpr.append(u)
    new_run.append(rpr)
    t = OxmlElement("w:t")
    t.text = text
    new_run.append(t)
    hyperlink.append(new_run)
    paragraph._p.append(hyperlink)


def add_inline(paragraph, text):
    for tok in INLINE.split(text):
        if not tok:
            continue
        if tok.startswith("**") and tok.endswith("**"):
            run = paragraph.add_run(tok[2:-2])
            run.bold = True
        elif tok.startswith("`") and tok.endswith("`"):
            run = paragraph.add_run(tok[1:-1])
            run.font.name = "Consolas"
            run.font.size = Pt(10)
        else:
            m = LINK.fullmatch(tok)
            if m:
                add_hyperlink(paragraph, m.group(2), m.group(1))
            else:
                paragraph.add_run(tok)


def strip_inline(text):
    text = LINK.sub(r"\1", text)
    text = text.replace("**", "").replace("`", "")
    return text


def is_table_sep(line):
    return bool(re.fullmatch(r"\s*\|?[\s:|-]+\|?\s*", line)) and "-" in line


def cells(line):
    line = line.strip()
    if line.startswith("|"):
        line = line[1:]
    if line.endswith("|"):
        line = line[:-1]
    return [c.strip() for c in line.split("|")]


def main():
    src, dst = sys.argv[1], sys.argv[2]
    with open(src, encoding="utf-8") as f:
        lines = f.read().split("\n")

    doc = Document()
    doc.styles["Normal"].font.name = "Calibri"
    doc.styles["Normal"].font.size = Pt(11)

    i = 0
    n = len(lines)
    while i < n:
        line = lines[i]
        stripped = line.strip()

        # fenced code
        if stripped.startswith("```"):
            i += 1
            buf = []
            while i < n and not lines[i].strip().startswith("```"):
                buf.append(lines[i])
                i += 1
            i += 1
            p = doc.add_paragraph()
            run = p.add_run("\n".join(buf))
            run.font.name = "Consolas"
            run.font.size = Pt(9)
            continue

        # table
        if "|" in line and i + 1 < n and is_table_sep(lines[i + 1]):
            header = cells(line)
            i += 2
            rows = []
            while i < n and "|" in lines[i] and lines[i].strip():
                rows.append(cells(lines[i]))
                i += 1
            table = doc.add_table(rows=1, cols=len(header))
            table.style = "Light Grid Accent 1"
            for j, h in enumerate(header):
                cell = table.rows[0].cells[j]
                cell.paragraphs[0].text = ""
                add_inline(cell.paragraphs[0], h)
                for r in cell.paragraphs[0].runs:
                    r.bold = True
            for row in rows:
                tr = table.add_row().cells
                for j in range(len(header)):
                    val = row[j] if j < len(row) else ""
                    tr[j].paragraphs[0].text = ""
                    add_inline(tr[j].paragraphs[0], val)
            doc.add_paragraph()
            continue

        # horizontal rule
        if re.fullmatch(r"\s*---+\s*", line):
            i += 1
            continue

        # headings
        m = re.match(r"^(#{1,6})\s+(.*)$", line)
        if m:
            level = len(m.group(1))
            doc.add_heading(strip_inline(m.group(2)), level=min(level, 4))
            i += 1
            continue

        # bullets
        m = re.match(r"^(\s*)[-*]\s+(.*)$", line)
        if m:
            p = doc.add_paragraph(style="List Bullet")
            add_inline(p, m.group(2))
            i += 1
            continue

        # numbered
        m = re.match(r"^(\s*)\d+\.\s+(.*)$", line)
        if m:
            p = doc.add_paragraph(style="List Number")
            add_inline(p, m.group(2))
            i += 1
            continue

        # blank
        if not stripped:
            i += 1
            continue

        # paragraph
        p = doc.add_paragraph()
        add_inline(p, stripped)
        i += 1

    doc.save(dst)
    print(f"Wrote {dst}")


if __name__ == "__main__":
    main()
